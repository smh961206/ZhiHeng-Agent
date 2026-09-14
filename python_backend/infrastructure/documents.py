from __future__ import annotations

import asyncio
import base64
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from functools import partial
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Protocol, cast

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from ..application.prompt_governance import build_prompt_plan
from ..domain.contracts import ApiError
from ..domain.research_budget import account as account_budget


class VisionGateway(Protocol):
    async def complete(self, purpose: str, messages: list[dict[str, Any]], *, stream: bool = False, max_tokens: int = 8000,
                       prompt_context: dict[str, Any] | None = None) -> str | Any: ...


async def _recognize_page(gateway: VisionGateway, page_content: list[dict[str, Any]]) -> Any:
    prompt = build_prompt_plan("vision-page-transcriber", [{"role": "user", "content": page_content}])
    return await gateway.complete(
        "vision", prompt.messages, stream=False, max_tokens=8000, prompt_context=prompt.telemetry
    )


class DocumentNeedsVision(ApiError):
    def __init__(self) -> None:
        super().__init__(400, "文件没有可读取正文；扫描件需要配置视觉模型")


def read_document(data: bytes, name: str) -> dict:
    if len(data) > 10 * 1024 * 1024:
        raise ApiError(413, "文件超过 10 MB")
    suffix = Path(name).suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("gb18030")
    elif suffix == ".pdf":
        reader = PdfReader(BytesIO(data), strict=True)
        if len(reader.pages) > 100:
            raise ApiError(400, "PDF最多100页")
        pages = [page.extract_text() or "" for page in reader.pages]
        if not any(page.strip() for page in pages):
            raise DocumentNeedsVision()
        text = "\n\n".join(f"【PDF第{i + 1}页】\n{page}" for i, page in enumerate(pages))
    elif suffix == ".docx":
        if len(zipfile.ZipFile(BytesIO(data)).infolist()) > 1200:
            raise ApiError(400, "Office文件结构超过限制")
        document = Document(BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
    elif suffix == ".xlsx":
        book = load_workbook(BytesIO(data), read_only=True, data_only=False)
        values = []
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows():
                cells = [{"address": cell.coordinate, "value": cell.value, "dataType": cell.data_type} for cell in row if cell.value is not None]
                if cells:
                    rows.append(cells)
            values.append({"sheet": sheet.title, "rows": rows})
        text = json.dumps(values, ensure_ascii=False, default=str)
    else:
        raise ApiError(400, "不支持的文件格式")
    text = re.sub(r"\x00", "", text).strip()
    if not text:
        if suffix == ".pdf":
            raise DocumentNeedsVision()
        raise ApiError(400, "文件没有可读取正文")
    return {"text": text[:20_000], "truncated": len(text) > 20_000, "parserVersion": "python-evidence-1", "documentPipeline": True}


def render_pdf_pages(data: bytes, *, max_pages: int = 8) -> list[bytes]:
    executable = shutil.which("pdftoppm")
    if not executable:
        raise ApiError(503, "扫描件识别组件不可用")
    reader = PdfReader(BytesIO(data), strict=True)
    if len(reader.pages) > 100:
        raise ApiError(400, "PDF最多100页")
    count = min(len(reader.pages), max_pages)
    images: list[bytes] = []
    with tempfile.TemporaryDirectory(prefix="zhiheng-vision-") as directory:
        source = Path(directory) / "source.pdf"
        source.write_bytes(data)
        for page in range(1, count + 1):
            target = Path(directory) / f"page-{page}"
            try:
                subprocess.run(
                    [executable, "-f", str(page), "-l", str(page), "-r", "120", "-png", "-singlefile", str(source), str(target)],
                    check=True,
                    timeout=30,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
            except (OSError, subprocess.SubprocessError) as error:
                raise ApiError(503, "扫描件页面渲染失败") from error
            image = target.with_suffix(".png").read_bytes()
            if len(image) > 8 * 1024 * 1024:
                raise ApiError(413, "扫描件单页图像超过识别上限")
            images.append(image)
    return images


async def read_document_with_vision(
    data: bytes,
    name: str,
    gateway: VisionGateway,
    *,
    renderer: Callable[[bytes], list[bytes]] = render_pdf_pages,
) -> dict[str, Any]:
    try:
        return await asyncio.to_thread(read_document, data, name)
    except DocumentNeedsVision:
        pages = await asyncio.to_thread(renderer, data)
        if not pages:
            raise ApiError(400, "扫描件没有可识别页面") from None
        transcripts: list[str] = []
        blocks: list[dict[str, Any]] = []
        recognized: dict[str, str] = {}
        for index, image in enumerate(pages, start=1):
            image_sha256 = hashlib.sha256(image).hexdigest()
            content = [
                {
                    "type": "text",
                    "text": "逐字转录这一页可见文字与表格。保留负号、括号、百分号、币种、单位和空值；不可猜测模糊数字，不作分析。模糊处写【无法辨认】。",
                },
                {"type": "image_url", "image_url": {"url": "data:image/png;base64," + base64.b64encode(image).decode("ascii")}},
            ]
            text = recognized.get(image_sha256, "")
            if not text:
                recognize = partial(_recognize_page, gateway, cast(list[dict[str, Any]], content))
                value = await account_budget("visionPage", 1, recognize)
                if not isinstance(value, str) or not value.strip():
                    raise ApiError(502, "视觉模型未返回有效转录") from None
                text = value.strip()
                recognized[image_sha256] = text
            transcripts.append(f"【PDF第{index}页·视觉转录待复核】\n{text}")
            blocks.append(
                {
                    "id": f"vision-b{index}",
                    "kind": "page-transcript",
                    "method": "vision",
                    "page": index,
                    "text": text,
                    "needsReview": True,
                    "imageSha256": image_sha256,
                    "instructionVersion": 1,
                }
            )
        joined = "\n\n".join(transcripts)
        return {
            "text": joined[:20_000],
            "truncated": len(joined) > 20_000 or len(PdfReader(BytesIO(data), strict=True).pages) > len(pages),
            "parserVersion": "python-evidence-vision-1",
            "documentPipeline": True,
            "documentBlocks": blocks,
            "method": "vision",
            "needsReview": True,
            "notice": "视觉转录只用于定位与人工复核，不能直接成为已验证财务事实。",
        }
